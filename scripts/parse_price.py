#!/usr/bin/env python3
"""Парсит прайс korshop (xlsx) -> src/data/products.json + public/img/*.webp"""
import json, re, sys, hashlib
from pathlib import Path
import openpyxl
from PIL import Image
import io

SRC = sys.argv[1] if len(sys.argv) > 1 else str(Path.home() / "Desktop/прайс 8.09.xlsx")
ROOT = Path(__file__).resolve().parent.parent
IMG_DIR = ROOT / "public/img"
IMG_DIR.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(SRC)
ws = wb[wb.sheetnames[0]]

# картинки по строкам (anchor 0-based -> excel row = +1)
images = {}
for im in ws._images:
    a = im.anchor._from
    if a.col != 2:      # картинки товаров лежат в колонке C
        continue
    images.setdefault(a.row + 1, im)

def num(v):
    if v is None: return None
    s = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try: return float(s)
    except ValueError: return None

def pack_qty(s):
    if not s: return None
    m = re.search(r"(\d+)\s*шт", str(s))
    return int(m.group(1)) if m else None

def category(name):
    n = name.upper()
    rules = [
        ("Наборы", ["НАБОР"]),
        ("Гаджеты", ["ГАДЖЕТ", "МАССАЖН", "РОЛЛЕР"]),
        ("Губы", ["ПОМАД", "ТИНТ", "БЛЕСК", "ПЛАМПЕР", "БАЛЬЗАМ ДЛЯ ГУБ", "ГУБ"]),
        ("Макияж", ["КУШОН", "ТУШЬ", "ТОНАЛЬН", "РУМЯН", "ПАЛЕТКА", "КОНСИЛЕР", "ПУДРА", "ТЕНИ", "КАРАНДАШ", "ХАЙЛАЙТЕР", "БАЗА ПОД", "ФИКСАТОР"]),
        ("Солнцезащита", ["СОЛНЦЕЗАЩИТ", "SPF", "SUNSCREEN"]),
        ("Маски и патчи", ["МАСК", "ПАТЧ", "ПЛАСТЫР"]),
        ("Очищение", ["ПЕНКА", "ГЕЛЬ ДЛЯ УМЫВАНИЯ", "ОЧИЩЕНИ", "МИЦЕЛЛЯР", "СКРАБ", "ГИДРОФИЛЬН", "ПИЛИНГ", "САЛФЕТКИ", "МОЛОЧКО", "СРЕДСТВО ДЛЯ УМЫВАНИЯ"]),
        ("Волосы", ["ШАМПУН", "ВОЛОС", "КОНДИЦИОНЕР"]),
        ("Тело", ["ТЕЛА", "ДЕЗОДОРАНТ", "МЫЛО", "ГЕЛЬ ДЛЯ ДУША", "РУК", "НОГ"]),
        ("Гигиена", ["ЗУБН", "ОПОЛАСКИВАТЕЛЬ"]),
        ("БАДы", ["БАД", "ВИТАМИН", "НАПИТОК", "ПОРОШОК"]),
        ("Уход за лицом", ["СЫВОРОТКА", "КРЕМ", "ЭССЕНЦИЯ", "ТОНЕР", "ПЭДЫ", "ЭМУЛЬСИЯ", "МИСТ", "СКВАЛАН", "АМПУЛ", "БУСТЕР", "ЛОСЬОН", "БАЛЬЗАМ", "СТИК", "ГЕЛЬ"]),
    ]
    for cat, keys in rules:
        if any(k in n for k in keys): return cat
    return "Прочее"

products, brand = [], None
sale = None
for i, row in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
    name, spec, _f, unit, barcode, pack, price, *_ = row
    if not name: continue
    name = str(name).strip()
    p = num(price)
    if p is None:                      # строка-заголовок бренда/раздела
        if re.match(r"^АКЦИЯ", name, re.I):
            sale, brand = name.strip(), name.strip()
        else:
            brand, sale = name, None
        continue
    im = images.get(i)
    img_name = None
    if im:
        data = im._data()
        h = hashlib.md5(data).hexdigest()[:12]
        img_name = f"{h}.webp"
        out = IMG_DIR / img_name
        if not out.exists():
            Image.open(io.BytesIO(data)).convert("RGB").save(out, "WEBP", quality=82, method=5)
    m = re.search(r"\(([^()]+)\)\s*(?:EXP.*)?$", name)
    real_brand = m.group(1).strip() if m else (brand or "")
    exp = re.search(r"EXP\s*([\d.]+)", name)
    clean = re.sub(r"\s*\(([^()]+)\)\s*(EXP\s*[\d.]+)?\s*$", "", name).strip()
    clean = re.sub(r"^АКЦИЯ\s+", "", clean)
    products.append({
        "id": len(products) + 1,
        "name": clean or name,
        "full": name,
        "brand": real_brand.upper(),
        "spec": (str(spec).strip() if spec else None),
        "unit": unit or "шт",
        "barcode": str(barcode) if barcode else None,
        "pack": (str(pack).replace("\n", " / ") if pack else None),
        "packQty": pack_qty(pack),
        "price": round(p, 2),
        "img": img_name,
        "cat": category(name),
        "sale": sale,
        "exp": exp.group(1) if exp else None,
    })

meta = {"currency": "USD", "date": "7 сентября 2026", "site": "www.korshop.one", "contact": "+996-559-050-618"}
out = ROOT / "src/data"
out.mkdir(parents=True, exist_ok=True)
(out / "products.json").write_text(json.dumps({"meta": meta, "products": products}, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"✅ {len(products)} товаров, {len(set(p['brand'] for p in products))} брендов, картинок: {sum(1 for p in products if p['img'])}")
